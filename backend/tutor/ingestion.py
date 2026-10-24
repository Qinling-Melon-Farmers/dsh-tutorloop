"""Read small local files and split into citeable source chunks."""
from __future__ import annotations
import hashlib
import io
import zipfile
from html.parser import HTMLParser
from pathlib import Path

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
MAX_PDF_PAGES = 500


class _ArticleParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.title_parts = [], []
        self.hidden = 0
        self.in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "svg", "template"}:
            self.hidden += 1
        if tag == "title":
            self.in_title = True
        if tag in {"p", "div", "article", "section", "br", "li", "h1", "h2", "h3", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "svg", "template"}:
            self.hidden = max(0, self.hidden - 1)
        if tag == "title":
            self.in_title = False
        if tag in {"p", "div", "article", "section", "li", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if self.in_title:
            self.title_parts.append(data)
        elif not self.hidden:
            self.parts.append(data)


def extract_html(html: str) -> tuple[str, str]:
    """Extract visible text without executing scripts or fetching resources."""
    parser = _ArticleParser()
    parser.feed(html)
    return "\n".join(s.strip() for s in "".join(parser.parts).splitlines() if s.strip()), "".join(parser.title_parts).strip()


def split_chunks(text: str, title: str, locator: str = "", url: str = "", kind: str = "local", size: int = 850, overlap: int = 100) -> list[dict]:
    if size <= overlap or overlap < 0 or size <= 0:
        raise ValueError("chunk size must be greater than overlap")
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    result = []
    step = size - overlap
    for i, start in enumerate(range(0, len(text), step)):
        piece = text[start:start + size].strip()
        if len(piece) < 30:
            continue
        loc = locator + (f" 第{i + 1}片" if len(text) > size else "")
        digest = hashlib.sha256(f"{title}|{loc}|{url}|{piece}".encode()).hexdigest()[:12]
        result.append({"id": "S" + digest, "title": title, "locator": loc,
                       "url": url, "kind": kind, "text": piece})
    return result


def parse_upload(filename: str, data: bytes) -> list[dict]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("文件超过20 MB限制，请按章节拆分后导入")
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError("PDF超过500页，请按章节拆分")
        parts = []
        total = 0
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            total += len(text)
            if total > MAX_TEXT_CHARS:
                raise ValueError("提取文本超过200万字符，请拆分文件")
            parts.extend(split_chunks(text, filename, f"第{i + 1}页"))
        if not parts:
            raise ValueError("PDF没有可提取正文，可能是扫描件；请先OCR后重新导入")
        return parts
    if suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError("DOCX解压体积超过80 MB，请拆分文件")
        from docx import Document
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(paragraphs)
        if len(text) > MAX_TEXT_CHARS:
            raise ValueError("文本超过200万字符，请拆分文件")
        return split_chunks(text, filename)
    if suffix in {".txt", ".md", ".html", ".htm"}:
        try:
            txt = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            txt = data.decode("gb18030", errors="replace")
        if len(txt) > MAX_TEXT_CHARS:
            raise ValueError("文本超过200万字符，请拆分文件")
        if suffix in {".html", ".htm"}:
            txt, _ = extract_html(txt)
        return split_chunks(txt, filename)
    raise ValueError("目前仅支持 PDF、DOCX、TXT、MD、HTML")


DEMO_TEXT = """某地老牌家具企业为降低成本，决定大量停用传统榫卯工艺，裁减经验丰富的木工，转而扩大自动化拼接家具生产。短期内产量提升、销售回暖，但部分新产品在使用后出现接缝松动与变形问题，售后投诉增加。老客户认为传统工艺是该品牌品质与文化特色的重要来源，不再继续购买。随着工匠离职，企业恢复传统工艺的能力也明显下降。
另一家企业选择用机械化生产普通产品，针对高端定制保留关键手工工序，并在生产后对品质进行统一检测。两种产品面向不同需求，企业会根据客户反馈调整生产比例。"""
